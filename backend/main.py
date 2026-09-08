from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any

from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .algorithms.costo_minimo import CostoMinimoSolver
from .algorithms.vogel import VogelSolver
from .models import ExportRequest, TransportationProblem

BASE_DIR = Path(__file__).resolve().parent.parent
FRONTEND_DIR = BASE_DIR / "frontend"

app = FastAPI(title="Problemas de Transporte", version="1.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.post("/costo-minimo")
def solve_costo_minimo(problem: TransportationProblem) -> dict:
    return _run_solver(CostoMinimoSolver, problem)


@app.post("/vogel")
def solve_vogel(problem: TransportationProblem) -> dict:
    return _run_solver(VogelSolver, problem)


@app.post("/comparar")
def compare_methods(problem: TransportationProblem) -> dict:
    min_cost_result = _run_solver(CostoMinimoSolver, problem)
    vogel_result = _run_solver(VogelSolver, problem)

    methods = [
        {"method": min_cost_result["method"], "total_cost": min_cost_result["total_cost"]},
        {"method": vogel_result["method"], "total_cost": vogel_result["total_cost"]},
    ]
    best_cost = min(item["total_cost"] for item in methods)

    return {
        "methods": [
            {
                **item,
                "is_best": item["total_cost"] == best_cost,
            }
            for item in methods
        ],
        "best_cost": best_cost,
        "results": {
            "costo_minimo": min_cost_result,
            "vogel": vogel_result,
        },
    }


@app.post("/export/pdf")
def export_pdf(payload: ExportRequest) -> StreamingResponse:
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    elements: list[Any] = [Paragraph(payload.title, styles["Title"]), Spacer(1, 12)]

    _append_result_to_pdf(elements, payload.result, styles)

    doc.build(elements)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": 'attachment; filename="resultados_transporte.pdf"'},
    )


@app.post("/export/excel")
def export_excel(payload: ExportRequest) -> StreamingResponse:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Resultados"

    _write_result_to_sheet(sheet, payload.result)

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="resultados_transporte.xlsx"'},
    )


def _run_solver(solver_class: type[CostoMinimoSolver] | type[VogelSolver], problem: TransportationProblem) -> dict:
    try:
        solver = solver_class(problem.costs, problem.supply, problem.demand)
        return solver.solve()
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


def _append_result_to_pdf(elements: list[Any], result: dict, styles: dict) -> None:
    if "results" in result:
        elements.append(Paragraph("Comparacion de metodos", styles["Heading2"]))
        comparison_data = [["Metodo", "Costo total"]]
        for item in result.get("methods", []):
            marker = " (menor costo)" if item.get("is_best") else ""
            comparison_data.append([f"{item.get('method')}{marker}", item.get("total_cost")])
        elements.append(_pdf_table(comparison_data))
        elements.append(Spacer(1, 12))

        for method_result in result["results"].values():
            _append_single_result_pdf(elements, method_result, styles)
    else:
        _append_single_result_pdf(elements, result, styles)


def _append_single_result_pdf(elements: list[Any], result: dict, styles: dict) -> None:
    elements.append(Paragraph(str(result.get("method", "Metodo")), styles["Heading2"]))
    elements.append(Paragraph(f"Costo total: {result.get('total_cost')}", styles["Normal"]))
    elements.append(Spacer(1, 8))

    elements.append(Paragraph("Matriz de costos balanceada", styles["Heading3"]))
    elements.append(_pdf_table(_matrix_table_data(result, "costs")))
    elements.append(Spacer(1, 8))

    elements.append(Paragraph("Matriz de asignacion final", styles["Heading3"]))
    elements.append(_pdf_table(_matrix_table_data(result, "allocations")))
    elements.append(Spacer(1, 8))

    fictitious_note = _fictitious_note(result)
    if fictitious_note:
        elements.append(Paragraph(fictitious_note, styles["Normal"]))
        elements.append(Spacer(1, 8))

    history_rows = [["Iteracion", "Decision", "Costo acumulado"]]
    for step in result.get("history", []):
        history_rows.append([step.get("iteration"), step.get("decision"), step.get("accumulated_cost")])
    elements.append(_pdf_table(history_rows))
    elements.append(Spacer(1, 14))


def _matrix_table_data(result: dict, key: str) -> list[list[Any]]:
    labels = result.get("destination_labels", [])
    origins = result.get("origin_labels", [])
    table_data = [[""] + labels]

    for index, row in enumerate(result.get(key, [])):
        label = origins[index] if index < len(origins) else f"O{index + 1}"
        table_data.append([label] + row)

    return table_data


def _fictitious_note(result: dict) -> str:
    balance = result.get("balance", {})
    if balance.get("balanced"):
        return "Oferta y demanda coinciden; no se agrego fila ni columna ficticia."
    if balance.get("added_type") == "column":
        return "La asignacion al destino ficticio representa inventario sobrante sin envio real."
    if balance.get("added_type") == "row":
        return "La asignacion al origen ficticio representa demanda insatisfecha dentro del balance del modelo."
    return ""


def _pdf_table(data: list[list[Any]]) -> Table:
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#172033")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#c8d0dc")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f6f8")]),
            ]
        )
    )
    return table


def _write_result_to_sheet(sheet: Any, result: dict) -> None:
    row_cursor = 1
    if "results" in result:
        sheet.cell(row=row_cursor, column=1, value="Comparacion de metodos")
        row_cursor += 1
        sheet.cell(row=row_cursor, column=1, value="Metodo")
        sheet.cell(row=row_cursor, column=2, value="Costo total")
        row_cursor += 1
        for item in result.get("methods", []):
            label = f"{item.get('method')} (menor costo)" if item.get("is_best") else item.get("method")
            sheet.cell(row=row_cursor, column=1, value=label)
            sheet.cell(row=row_cursor, column=2, value=item.get("total_cost"))
            row_cursor += 1
        row_cursor += 2

        for method_result in result["results"].values():
            row_cursor = _write_single_result_sheet(sheet, method_result, row_cursor)
            row_cursor += 2
    else:
        _write_single_result_sheet(sheet, result, row_cursor)

    for column_cells in sheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        sheet.column_dimensions[column_cells[0].column_letter].width = min(max_length + 2, 60)


def _write_single_result_sheet(sheet: Any, result: dict, row_cursor: int) -> int:
    sheet.cell(row=row_cursor, column=1, value=result.get("method", "Metodo"))
    row_cursor += 1
    sheet.cell(row=row_cursor, column=1, value="Costo total")
    sheet.cell(row=row_cursor, column=2, value=result.get("total_cost"))
    row_cursor += 2

    sheet.cell(row=row_cursor, column=1, value="Matriz de costos balanceada")
    row_cursor += 1
    row_cursor = _write_matrix_sheet(sheet, result, "costs", row_cursor)
    row_cursor += 1

    sheet.cell(row=row_cursor, column=1, value="Matriz de asignacion")
    row_cursor += 1
    row_cursor = _write_matrix_sheet(sheet, result, "allocations", row_cursor)

    row_cursor += 1
    note = _fictitious_note(result)
    if note:
        sheet.cell(row=row_cursor, column=1, value=note)
        row_cursor += 2

    sheet.cell(row=row_cursor, column=1, value="Historial")
    row_cursor += 1
    sheet.cell(row=row_cursor, column=1, value="Iteracion")
    sheet.cell(row=row_cursor, column=2, value="Decision")
    sheet.cell(row=row_cursor, column=3, value="Costo acumulado")
    row_cursor += 1

    for step in result.get("history", []):
        sheet.cell(row=row_cursor, column=1, value=step.get("iteration"))
        sheet.cell(row=row_cursor, column=2, value=step.get("decision"))
        sheet.cell(row=row_cursor, column=3, value=step.get("accumulated_cost"))
        row_cursor += 1

    return row_cursor


def _write_matrix_sheet(sheet: Any, result: dict, key: str, row_cursor: int) -> int:
    for col_index, label in enumerate(result.get("destination_labels", []), start=2):
        sheet.cell(row=row_cursor, column=col_index, value=label)
    row_cursor += 1

    for index, matrix_row in enumerate(result.get(key, [])):
        label = result.get("origin_labels", [])[index]
        sheet.cell(row=row_cursor, column=1, value=label)
        for col_index, value in enumerate(matrix_row, start=2):
            sheet.cell(row=row_cursor, column=col_index, value=value)
        row_cursor += 1

    return row_cursor


@app.get("/")
def index() -> FileResponse:
    return FileResponse(FRONTEND_DIR / "index.html")


app.mount("/static", StaticFiles(directory=FRONTEND_DIR), name="static")
